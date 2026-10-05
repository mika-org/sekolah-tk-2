# -*- coding: utf-8 -*-
"""
Manual Book Generator for Smart Kids - YAPCHI Multi-School Management System
Generates a complete, publication-ready Microsoft Word (.docx) user manual.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

# --- COLOR PALETTE ---
COLOR_PRIMARY = RGBColor(6, 95, 70)       # Deep Emerald (#065F46)
COLOR_SECONDARY = RGBColor(15, 118, 110)  # Deep Teal (#0F766E)
COLOR_TEXT_MAIN = RGBColor(30, 41, 59)    # Slate 800 (#1E293B)
COLOR_TEXT_MUTED = RGBColor(100, 116, 139)# Slate 500 (#64748B)
COLOR_HIGHLIGHT = RGBColor(4, 120, 87)    # Emerald 700 (#047857)

HEX_PRIMARY = "065F46"
HEX_SECONDARY = "0F766E"
HEX_BG_LIGHT = "F8FAFC"
HEX_BG_CALLOUT_INFO = "ECFDF5"
HEX_BG_CALLOUT_WARN = "FFFBEB"
HEX_BG_CALLOUT_DANGER = "FEF2F2"
HEX_BORDER_TABLE = "CBD5E1"
HEX_BORDER_IMG = "059669"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, top="none", bottom="none", left="none", right="none", 
                     color="CBD5E1", sz="8"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders_elm = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders_elm)

class ManualBookBuilder:
    def __init__(self):
        self.doc = docx.Document()
        self.setup_page_properties()
        self.img_counter = 0

    def setup_page_properties(self):
        section = self.doc.sections[0]
        # A4 paper size
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Buku Panduan Penggunaan Sistem Smart Kids | YAPCHI Foundation")
        hrun.font.name = "Segoe UI"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = COLOR_TEXT_MUTED

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun1 = fp.add_run("Halaman ")
        frun1.font.name = "Segoe UI"
        frun1.font.size = Pt(9)
        frun1.font.color.rgb = COLOR_TEXT_MUTED
        
        fld = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
        fp._p.append(fld)

    def add_cover_page(self):
        # Top space
        for _ in range(2):
            self.doc.add_paragraph()

        # Badge pill
        tbl = self.doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(4.5)
        set_cell_background(cell, "ECFDF5")
        set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
        set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="059669", sz="6")
        bp = cell.paragraphs[0]
        bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        brun = bp.add_run("DOKUMEN RESMI PANDUAN PENGGUNA (USER MANUAL)")
        brun.font.name = "Segoe UI"
        brun.font.bold = True
        brun.font.size = Pt(10)
        brun.font.color.rgb = RGBColor(4, 120, 87)

        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(24)
        sp.paragraph_format.space_after = Pt(8)
        sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        trun = sp.add_run("SISTEM INFORMASI MANAJEMEN\nMULTI-SEKOLAH & PPDB ONLINE")
        trun.font.name = "Segoe UI"
        trun.font.bold = True
        trun.font.size = Pt(24)
        trun.font.color.rgb = COLOR_PRIMARY

        sub_p = self.doc.add_paragraph()
        sub_p.paragraph_format.space_before = Pt(4)
        sub_p.paragraph_format.space_after = Pt(28)
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        srun = sub_p.add_run("Panduan Operasional Terlengkap untuk Administrator, Tenaga Pendidik (Guru),\nOrang Tua / Wali Murid, dan Pengunjung Publik")
        srun.font.name = "Segoe UI"
        srun.font.size = Pt(12)
        srun.font.color.rgb = COLOR_TEXT_MAIN

        # Decorative line
        dp = self.doc.add_paragraph()
        dp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        dp.paragraph_format.space_after = Pt(28)
        drun = dp.add_run("━━━━━━━━━━━━━━━━━━━ ◆ ━━━━━━━━━━━━━━━━━━━")
        drun.font.name = "Segoe UI"
        drun.font.size = Pt(12)
        drun.font.color.rgb = COLOR_SECONDARY

        # Metadata box
        tbl_meta = self.doc.add_table(rows=6, cols=2)
        tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_data = [
            ("Nama Platform", "Smart Kids Cloud Management System (CMS)"),
            ("Naungan Yayasan", "YAPCHI Foundation (Yayasan Pendidikan Anak Indonesia)"),
            ("Cabang Sekolah Terpadu", "Smart Kids Sadjati, Smart Kids BCL & Cabang Afiliasi"),
            ("Versi Rilis Dokumen", "Versi 2.0 (Edisi Lengkap Multi-Sekolah)"),
            ("Tahun Akademik", "Tahun Ajaran 2026 / 2027"),
            ("Status Dokumen", "Buku Panduan Standar Operasional Prosedur (SOP)"),
        ]
        for row_idx, (k, v) in enumerate(meta_data):
            cell_k = tbl_meta.cell(row_idx, 0)
            cell_v = tbl_meta.cell(row_idx, 1)
            cell_k.width = Inches(2.2)
            cell_v.width = Inches(4.0)
            set_cell_background(cell_k, "F1F5F9" if row_idx % 2 == 0 else "FFFFFF")
            set_cell_background(cell_v, "F8FAFC" if row_idx % 2 == 0 else "FFFFFF")
            set_cell_margins(cell_k, top=100, bottom=100, left=140, right=140)
            set_cell_margins(cell_v, top=100, bottom=100, left=140, right=140)
            set_cell_borders(cell_k, top="single", bottom="single", left="none", right="none", color="CBD5E1", sz="4")
            set_cell_borders(cell_v, top="single", bottom="single", left="none", right="none", color="CBD5E1", sz="4")

            pk = cell_k.paragraphs[0]
            pk.paragraph_format.space_after = Pt(0)
            rk = pk.add_run(k)
            rk.font.name = "Segoe UI"
            rk.font.bold = True
            rk.font.size = Pt(10)
            rk.font.color.rgb = COLOR_TEXT_MAIN

            pv = cell_v.paragraphs[0]
            pv.paragraph_format.space_after = Pt(0)
            rv = pv.add_run(v)
            rv.font.name = "Segoe UI"
            rv.font.size = Pt(10)
            rv.font.color.rgb = COLOR_TEXT_MAIN

        # Bottom notice
        for _ in range(3):
            self.doc.add_paragraph()

        bp2 = self.doc.add_paragraph()
        bp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_bot = bp2.add_run("Pusat Layanan & Pengembangan Teknologi Pendidikan — YAPCHI Foundation\nKarawang, Jawa Barat - Indonesia")
        r_bot.font.name = "Segoe UI"
        r_bot.font.size = Pt(9.5)
        r_bot.font.color.rgb = COLOR_TEXT_MUTED

        self.doc.add_page_break()

    def add_page_break(self):
        self.doc.add_page_break()

    def add_heading_1(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Segoe UI"
        run.font.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = COLOR_PRIMARY

        # Bottom horizontal accent bar
        tbl = self.doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.27)
        set_cell_background(cell, "059669")
        set_cell_margins(cell, top=10, bottom=10, left=0, right=0)
        set_cell_borders(cell, top="none", bottom="none", left="none", right="none")
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)

        # Spacing after bar
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(4)
        sp.paragraph_format.space_after = Pt(0)

    def add_heading_2(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Segoe UI"
        run.font.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = COLOR_SECONDARY

    def add_heading_3(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Segoe UI"
        run.font.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = COLOR_TEXT_MAIN

    def add_paragraph(self, text, bold_prefix=None, italic=False):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Segoe UI"
            r_pre.font.bold = True
            r_pre.font.size = Pt(10)
            r_pre.font.color.rgb = COLOR_TEXT_MAIN
        run = p.add_run(text)
        run.font.name = "Segoe UI"
        run.font.italic = italic
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_TEXT_MAIN
        return p

    def add_bullet(self, text, bold_prefix=None):
        p = self.doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Segoe UI"
            r_pre.font.bold = True
            r_pre.font.size = Pt(10)
            r_pre.font.color.rgb = COLOR_TEXT_MAIN
        run = p.add_run(text)
        run.font.name = "Segoe UI"
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_TEXT_MAIN
        return p

    def add_callout(self, text, title="CATATAN PENTING", callout_type="info"):
        tbl = self.doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.27)

        if callout_type == "info":
            bg = HEX_BG_CALLOUT_INFO
            border_col = "059669"
            t_col = RGBColor(4, 120, 87)
            icon = "ℹ️"
        elif callout_type == "warning":
            bg = HEX_BG_CALLOUT_WARN
            border_col = "D97706"
            t_col = RGBColor(180, 83, 9)
            icon = "⚠️"
        else:
            bg = HEX_BG_CALLOUT_DANGER
            border_col = "DC2626"
            t_col = RGBColor(220, 38, 38)
            icon = "🚨"

        set_cell_background(cell, bg)
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        set_cell_borders(cell, top="none", bottom="none", left="single", right="none", color=border_col, sz="24")

        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(3)
        r_head = p.add_run(f"{icon} {title}")
        r_head.font.name = "Segoe UI"
        r_head.font.bold = True
        r_head.font.size = Pt(10)
        r_head.font.color.rgb = t_col

        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.line_spacing = 1.15
        r_txt = p2.add_run(text)
        r_txt.font.name = "Segoe UI"
        r_txt.font.size = Pt(9.5)
        r_txt.font.color.rgb = COLOR_TEXT_MAIN

        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(4)
        sp.paragraph_format.space_after = Pt(0)

    def add_image_placeholder(self, title, description, route, tips):
        self.img_counter += 1
        num_str = f"{self.img_counter:02d}"

        tbl = self.doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.27)

        set_cell_background(cell, HEX_BG_LIGHT)
        set_cell_margins(cell, top=160, bottom=160, left=200, right=200)
        set_cell_borders(cell, top="dashed", bottom="dashed", left="dashed", right="dashed", color=HEX_BORDER_IMG, sz="12")

        # Header Title
        p0 = cell.paragraphs[0]
        p0.paragraph_format.space_after = Pt(4)
        r0 = p0.add_run(f"🖼️ [TEMPAT GAMBAR {num_str}: {title}]")
        r0.font.name = "Segoe UI"
        r0.font.bold = True
        r0.font.size = Pt(10.5)
        r0.font.color.rgb = COLOR_PRIMARY

        # Route / Path
        p1 = cell.add_paragraph()
        p1.paragraph_format.space_after = Pt(2)
        r1_lbl = p1.add_run("📍 Lokasi Halaman / Rute: ")
        r1_lbl.font.name = "Segoe UI"
        r1_lbl.font.bold = True
        r1_lbl.font.size = Pt(9)
        r1_lbl.font.color.rgb = COLOR_SECONDARY
        r1_val = p1.add_run(route)
        r1_val.font.name = "Consolas"
        r1_val.font.size = Pt(9)
        r1_val.font.color.rgb = COLOR_TEXT_MAIN

        # Description
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(2)
        r2_lbl = p2.add_run("📝 Deskripsi Tampilan: ")
        r2_lbl.font.name = "Segoe UI"
        r2_lbl.font.bold = True
        r2_lbl.font.size = Pt(9)
        r2_lbl.font.color.rgb = COLOR_TEXT_MAIN
        r2_val = p2.add_run(description)
        r2_val.font.name = "Segoe UI"
        r2_val.font.size = Pt(9)
        r2_val.font.color.rgb = COLOR_TEXT_MAIN

        # Tips
        p3 = cell.add_paragraph()
        p3.paragraph_format.space_after = Pt(8)
        r3_lbl = p3.add_run("💡 Panduan Pengambilan Screenshot: ")
        r3_lbl.font.name = "Segoe UI"
        r3_lbl.font.bold = True
        r3_lbl.font.size = Pt(8.5)
        r3_lbl.font.color.rgb = RGBColor(4, 120, 87)
        r3_val = p3.add_run(tips)
        r3_val.font.name = "Segoe UI"
        r3_val.font.italic = True
        r3_val.font.size = Pt(8.5)
        r3_val.font.color.rgb = RGBColor(4, 120, 87)

        # Dropzone indicator
        p4 = cell.add_paragraph()
        p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p4.paragraph_format.space_before = Pt(6)
        p4.paragraph_format.space_after = Pt(2)
        r4 = p4.add_run("▼ SILAKAN TEMPELKAN TANGKAPAN LAYAR (SCREENSHOT) DI BAWAH INI ATAU GANTIKAN KOTAK INI DENGAN GAMBAR ▼")
        r4.font.name = "Segoe UI"
        r4.font.bold = True
        r4.font.size = Pt(8.5)
        r4.font.color.rgb = RGBColor(148, 163, 184)

        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(6)
        sp.paragraph_format.space_after = Pt(0)

    def add_table(self, headers, rows, col_widths=None):
        tbl = self.doc.add_table(rows=len(rows) + 1, cols=len(headers))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        # Header Row
        for col_idx, h_text in enumerate(headers):
            cell = tbl.cell(0, col_idx)
            if col_widths and col_idx < len(col_widths):
                cell.width = Inches(col_widths[col_idx])
            set_cell_background(cell, HEX_PRIMARY)
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            set_cell_borders(cell, top="single", bottom="single", left="none", right="none", color="047857", sz="8")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(h_text)
            r.font.name = "Segoe UI"
            r.font.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(255, 255, 255)

        # Data Rows
        for row_idx, row_data in enumerate(rows):
            bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, cell_value in enumerate(row_data):
                cell = tbl.cell(row_idx + 1, col_idx)
                if col_widths and col_idx < len(col_widths):
                    cell.width = Inches(col_widths[col_idx])
                set_cell_background(cell, bg)
                set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
                set_cell_borders(cell, top="single", bottom="single", left="none", right="none", color=HEX_BORDER_TABLE, sz="4")
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                r = p.add_run(str(cell_value))
                r.font.name = "Segoe UI"
                r.font.size = Pt(9)
                r.font.color.rgb = COLOR_TEXT_MAIN

        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(4)
        sp.paragraph_format.space_after = Pt(0)

    def save(self, filepath):
        self.doc.save(filepath)
        print(f"Document successfully saved at: {filepath}")
        print(f"Total image placeholders generated: {self.img_counter}")

if __name__ == "__main__":
    builder = ManualBookBuilder()
    builder.add_cover_page()
    builder.add_heading_1("BAB 1: PENDAHULUAN & GAMBARAN UMUM SISTEM")
    builder.add_paragraph("Sistem Smart Kids CMS merupakan platform digital terpadu...")
    builder.add_image_placeholder(
        "Tampilan Halaman Utama Website Publik", 
        "Landing page utama menampilkan banner PPDB, logo, navbar cabang, dan video profil.",
        "URL: / (Home)", 
        "Buka browser pada halaman utama, ambil screenshot layar penuh bagian hero section."
    )
    builder.save("scratch/test_builder.docx")
