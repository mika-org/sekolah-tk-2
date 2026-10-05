import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def test_footer():
    doc = docx.Document()
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = "Buku Panduan Sistem Smart Kids | Halaman "
    run = p.add_run()
    
    fld1 = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
    p._p.append(fld1)
    
    doc.save("scratch/test_footer.docx")
    print("fldSimple PAGE works perfectly!")

if __name__ == "__main__":
    test_footer()
