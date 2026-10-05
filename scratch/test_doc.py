import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_test_doc():
    doc = docx.Document()
    
    # Set standard A4 page size
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    
    p = doc.add_paragraph()
    run = p.add_run("Testing Helpers")
    run.font.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(6, 95, 70)
    
    doc.save("scratch/test_out.docx")
    print("Test document created successfully")

if __name__ == "__main__":
    create_test_doc()
