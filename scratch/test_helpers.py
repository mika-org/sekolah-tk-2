import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_borders(cell, top="none", bottom="none", left="none", right="none", 
                     color="CBD5E1", sz="8"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders_elm = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders_elm)

def test_helpers():
    doc = docx.Document()
    
    # Test placeholder box
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.2)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=180, bottom=180, left=240, right=240)
    set_cell_borders(cell, top="dashed", bottom="dashed", left="dashed", right="dashed", color="059669", sz="12")
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("📷 [TEMPAT GAMBAR 01: Form Login Multi-Role]")
    r.font.name = "Segoe UI"
    r.font.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(4, 120, 87)
    
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run("Deskripsi: ")
    r2.font.bold = True
    r2.font.size = Pt(9.5)
    r2.font.name = "Segoe UI"
    r2_txt = p2.add_run("Tampilan form login pada rute /admin/login")
    r2_txt.font.size = Pt(9.5)
    r2_txt.font.name = "Segoe UI"
    
    doc.save("scratch/test_helpers_out.docx")
    print("test_helpers passed")

if __name__ == "__main__":
    test_helpers()
